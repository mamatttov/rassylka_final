"""Чтение зарплатной ведомости.

Понимает два формата:
1. Простая таблица (CSV или Excel): Таб№, ФИО, Email, ... и столбцы
   «Начислено: Оклад», «Удержано: ПН», «Выплачено: Аванс» и т.п.
2. «Расчетная ведомость» из 1С как она выгружается в Excel
   (объединённые заголовки, подразделения отдельными строками, без email).
"""
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from utils import MONTHS, money, read_table, simple, to_kopeyki

SECTIONS = ["Начислено", "Удержано", "Выплачено"]


@dataclass
class Sotrudnik:
    fio: str
    tab: str = ""
    email: str = ""
    position: str = ""
    dept: str = ""
    days: str = ""
    lines: dict = field(default_factory=lambda: {s: [] for s in SECTIONS})  # раздел -> [(статья, тыйыны)]
    saldo_start: int = 0       # долг на начало месяца (бывает в 1С)
    k_vyplate_1c: int = None   # «Сальдо на конец» из 1С, для сверки
    skip: str = ""             # причина, по которой письмо не отправляем
    row: int = 0               # номер строки в файле, чтобы было где искать ошибку

    @property
    def key(self):
        # по этому ключу ведётся журнал отправок; в 1С табельного номера нет — берём ФИО
        return self.tab or self.fio

    def total(self, section):
        return sum(kop for _, kop in self.lines[section])

    @property
    def k_vyplate(self):
        return self.saldo_start + self.total("Начислено") - self.total("Удержано") - self.total("Выплачено")


def read_vedomost(path):
    """Возвращает (сотрудники, замечания, {'org': ..., 'period': ...})."""
    if Path(path).suffix.lower() in (".xlsx", ".xlsm"):
        sheet = load_workbook(path, data_only=True).worksheets[0]
        header_row = find_1c_header(sheet)
        if header_row:
            return read_1c(sheet, header_row)

    people, problems = read_simple(read_table(path))
    return people, problems, {}


def check(people, problems):
    """Общие проверки для обоих форматов."""
    for p in people:
        if p.k_vyplate < 0:
            problems.append(f"строка {p.row} ({p.fio}): к выплате получается минус {money(-p.k_vyplate)} — проверьте суммы")
        if p.k_vyplate_1c is not None and p.k_vyplate_1c != p.k_vyplate:
            problems.append(f"строка {p.row} ({p.fio}): в ведомости к выплате {money(p.k_vyplate_1c)}, "
                            f"а по расчёту {money(p.k_vyplate)}")
    tabs = [p.tab for p in people if p.tab]
    for tab in sorted(set(tabs)):
        if tabs.count(tab) > 1:
            problems.append(f"табельный номер {tab} встречается несколько раз")


# ---------- простая таблица ----------

def read_simple(rows):
    if not rows:
        raise SystemExit("Ведомость пустая.")

    columns = list(rows[0].keys())

    def find_column(*names):
        for col in columns:
            if simple(col) in [simple(n) for n in names]:
                return col
        return None

    col_tab = find_column("Таб№", "Табельный номер", "Таб")
    col_fio = find_column("ФИО", "Сотрудник")
    col_email = find_column("Email", "E-mail", "Почта")
    col_position = find_column("Должность")
    col_dept = find_column("Подразделение", "Отдел")
    col_days = find_column("Отработано дней", "Дней")
    col_itog = find_column("К выплате", "Итого к выплате")

    if not col_fio or not col_email:
        raise SystemExit(f"В ведомости нужны столбцы «ФИО» и «Email». Сейчас есть: {', '.join(columns)}")

    # столбцы сумм: «Начислено: Оклад» -> раздел «Начислено», статья «Оклад»
    money_columns = []
    for col in columns:
        if ":" not in col:
            continue
        section, item = col.split(":", 1)
        section = section.strip().capitalize()
        if section in SECTIONS:
            money_columns.append((col, section, item.strip()))
    if not any(section == "Начислено" for _, section, _ in money_columns):
        raise SystemExit("В ведомости нет ни одного столбца «Начислено: ...».")

    people, problems = [], []
    for row_number, row in enumerate(rows, start=2):
        fio = (row.get(col_fio) or "").strip()
        if not fio:
            continue
        p = Sotrudnik(fio=fio, row=row_number)
        p.tab = (row.get(col_tab) or "").strip() if col_tab else ""
        p.email = (row.get(col_email) or "").strip()
        p.position = (row.get(col_position) or "").strip() if col_position else ""
        p.dept = (row.get(col_dept) or "").strip() if col_dept else ""
        p.days = (row.get(col_days) or "").strip() if col_days else ""

        try:
            for col, section, item in money_columns:
                kop = to_kopeyki(row.get(col))
                if kop:
                    p.lines[section].append((item, kop))
            if col_itog and (row.get(col_itog) or "").strip():
                p.k_vyplate_1c = to_kopeyki(row[col_itog])
        except Exception:
            problems.append(f"строка {row_number} ({fio}): не получилось прочитать сумму")
            continue

        if not p.email:
            p.skip = "нет email"
        elif not looks_like_email(p.email):
            p.skip = f"неверный email «{p.email}»"
        people.append(p)

    check(people, problems)
    return people, problems


def looks_like_email(text):
    return re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", text) is not None


# ---------- ведомость из 1С ----------

# Как называются статьи в 1С и как их понятнее показать сотруднику
NICE_NAMES = {
    "пфр": "Страховые взносы в Соцфонд (ПФР)",
    "гнпфр": "Накопительный пенсионный фонд (ГНПФ)",
    "пн": "Подоходный налог",
}


def find_1c_header(sheet):
    """Ищем строку, где есть «Начислено», а под ней в первом столбце — «Сотрудник»."""
    for r in range(1, min(sheet.max_row, 30)):
        titles = [simple(c.value) for c in sheet[r] if c.value is not None]
        below = simple(sheet.cell(r + 1, 1).value or "")
        if "начислено" in titles and below == "сотрудник":
            return r
    return None


def read_1c(sheet, header_row):
    info = {"org": str(sheet.cell(1, 1).value or "").strip()}

    # «Расчетная ведомость за Октябрь 2026 г.» -> 2026-10
    title = str(sheet.cell(2, 1).value or "")
    found = re.search(r"за\s+(\w+)\s+(\d{4})", title, re.IGNORECASE)
    if found and found.group(1).lower() in MONTHS:
        month = MONTHS.index(found.group(1).lower()) + 1
        info["period"] = f"{found.group(2)}-{month:02d}"

    # Разбираем шапку. В 1С после «Начислено» идут его статьи, после «Удержано» — свои,
    # потом «Удержано с фирмы» (взносы работодателя — сотруднику их не показываем).
    group_total = {}     # раздел -> номер столбца с итогом
    group_items = {}     # раздел -> [(номер столбца, статья)]
    other = {}           # дней / сальдо / выплачено -> номер столбца
    current = None
    for col in range(2, sheet.max_column + 1):
        name = str(sheet.cell(header_row, col).value or "").strip()
        key = simple(name)
        if not key:
            continue
        if key == "начислено":
            current = "Начислено"
        elif key == "удержано":
            current = "Удержано"
        elif key == "удержаносфирмы":
            current = "с фирмы"
        elif key in ("дней", "часов", "сальдонаначало", "выплачено", "сальдонаконец"):
            other[key] = col
            current = None
            continue
        elif current:
            group_items[current].append((col, NICE_NAMES.get(key, name.replace("\n", " "))))
            continue
        else:
            continue
        group_total[current] = col
        group_items[current] = []

    sub_header = {simple(sheet.cell(header_row + 1, c).value or ""): c for c in range(1, sheet.max_column + 1)}
    col_position = sub_header.get("должность")
    money_cols = list(group_total.values()) + [c for k, c in other.items() if k not in ("дней", "часов")]

    def amount(row, col):
        return to_kopeyki(sheet.cell(row, col).value)

    people, problems = [], []
    dept = ""
    for r in range(header_row + 2, sheet.max_row + 1):
        name = str(sheet.cell(r, 1).value or "").strip()
        if not name:
            continue
        if name.lower().startswith("итого"):
            break
        if all(amount(r, c) == 0 for c in money_cols):
            dept = name  # строка без сумм — это название подразделения
            continue

        p = Sotrudnik(fio=name, dept=dept, row=r)
        if col_position:
            p.position = str(sheet.cell(r, col_position).value or "").strip()
        if "дней" in other:
            p.days = str(sheet.cell(r, other["дней"]).value or "")

        try:
            for section in ("Начислено", "Удержано"):
                if section not in group_total:
                    continue
                for col, item in group_items[section]:
                    kop = amount(r, col)
                    if kop:
                        p.lines[section].append((item, kop))
                # если статьи не сходятся с итогом — разницу показываем отдельной строкой
                diff = amount(r, group_total[section]) - p.total(section)
                if diff:
                    p.lines[section].append(("Прочее", diff))

            if "выплачено" in other and amount(r, other["выплачено"]):
                p.lines["Выплачено"].append(("Выплачено в течение месяца", amount(r, other["выплачено"])))
            if "сальдонаначало" in other:
                p.saldo_start = amount(r, other["сальдонаначало"])
            if "сальдонаконец" in other:
                p.k_vyplate_1c = amount(r, other["сальдонаконец"])
        except Exception:
            problems.append(f"строка {r} ({name}): не получилось прочитать сумму")
            continue

        people.append(p)

    check(people, problems)
    return people, problems, info
