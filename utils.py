"""Общие мелочи: суммы, даты, чтение таблиц."""
import csv
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

MONTHS = ["январь", "февраль", "март", "апрель", "май", "июнь",
          "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]


def to_kopeyki(value):
    """Сумму из ячейки ('1 234,50', 1234.5, None) переводим в целые тыйыны.

    Считаем в целых числах, чтобы не было ошибок округления float.
    """
    if value is None:
        return 0
    text = str(value).strip()
    if text == "" or text.startswith("="):  # пустая ячейка или формула без значения
        return 0
    text = text.replace("\xa0", "").replace(" ", "").replace(",", ".")
    kop = Decimal(text) * 100
    return int(kop.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def money(kop):
    """47900 сом в тыйынах -> '47 900.00'"""
    rub = Decimal(kop) / 100
    return f"{rub:,.2f}".replace(",", " ")


def month_title(period):
    """'2026-10' -> 'октябрь 2026'"""
    year, month = period.split("-")
    return f"{MONTHS[int(month) - 1]} {year}"


def simple(text):
    """Название столбца без пробелов, точек и регистра: 'Таб. №' -> 'таб'"""
    return re.sub(r"[\s.№#]", "", str(text).lower())


def simple_fio(fio):
    """Чтобы 'Иванов  Пётр' и 'иванов петр' считались одним человеком."""
    return " ".join(fio.lower().replace("ё", "е").split())


def read_csv_file(path):
    # Excel сохраняет CSV то в UTF-8, то в Windows-1251, а разделитель бывает ; или ,
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            with open(path, encoding=encoding, newline="") as f:
                sample = f.read(2048)
                f.seek(0)
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
                return list(csv.DictReader(f, dialect=dialect))
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"Не получилось прочитать {path}: неизвестная кодировка.")


def read_table(path):
    """CSV или первый лист Excel -> список словарей {заголовок: значение}."""
    path = Path(path)
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return read_csv_file(path)

    sheet = load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        if all(v in (None, "") for v in row):
            continue
        item = {}
        for name, value in zip(header, row):
            item[name] = "" if value is None else str(value)
        result.append(item)
    return result
