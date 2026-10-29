"""Расчётный листок сотрудника в Excel."""
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from utils import month_title

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="DDE7F3")
MONEY_FORMAT = "#,##0.00"

SECTION_TITLES = {"Начислено": "Начислено", "Удержано": "Удержано", "Выплачено": "Выплачено ранее"}


def make_listok(p, period, org, folder):
    wb = Workbook()
    ws = wb.active
    ws.title = "Расчётный листок"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 18

    ws["A1"] = org
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = f"Расчётный листок за {month_title(period)}"
    ws["A2"].font = Font(bold=True, size=13)

    row = 4
    for label, value in [("Сотрудник", p.fio), ("Табельный №", p.tab), ("Должность", p.position),
                         ("Подразделение", p.dept), ("Отработано дней", p.days)]:
        if not value:
            continue
        ws.cell(row, 1, label).font = Font(color="555555")
        ws.cell(row, 2, value).alignment = Alignment(horizontal="left")
        row += 1

    for section, title in SECTION_TITLES.items():
        if not p.lines[section]:
            continue
        row += 1
        write_header(ws, row, title)
        for item, kop in p.lines[section]:
            row += 1
            ws.cell(row, 1, item).border = BORDER
            write_money(ws, row, kop).border = BORDER
        row += 1
        ws.cell(row, 1, f"Итого {title.lower()}").font = Font(bold=True)
        write_money(ws, row, p.total(section)).font = Font(bold=True)

    if p.saldo_start:
        row += 2
        if p.saldo_start > 0:
            ws.cell(row, 1, "Долг за предприятием на начало месяца")
        else:
            ws.cell(row, 1, "Долг работника на начало месяца")
        write_money(ws, row, abs(p.saldo_start))

    row += 2
    ws.cell(row, 1, "К ВЫПЛАТЕ").font = Font(bold=True, size=13)
    write_money(ws, row, p.k_vyplate).font = Font(bold=True, size=13)

    name = re.sub(r"[^\w-]+", "_", p.key)  # ФИО с пробелами -> Иванов_Иван
    path = Path(folder) / f"Расчетный_листок_{period}_{name}.xlsx"
    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"Не получается сохранить {path.name}: файл открыт в Excel. Закройте его и запустите снова.")
    return path


def write_header(ws, row, title):
    for col, text in [(1, title), (2, "Сумма")]:
        cell = ws.cell(row, col, text)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER


def write_money(ws, row, kop):
    cell = ws.cell(row, 2, kop / 100)
    cell.number_format = MONEY_FORMAT
    return cell
