"""Справочник сотрудников: ФИО -> email.

Нужен для ведомости из 1С, в которой нет почты.
Файл sotrudniki.xlsx (столбцы ФИО, Email, Примечание) правится руками в Excel.
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from listok import BORDER, HEADER_FILL
from utils import read_table, simple, simple_fio
from vedomost import looks_like_email

HERE = Path(__file__).parent
DEFAULT_FILES = [HERE / "sotrudniki.xlsx", HERE / "sotrudniki.csv"]


def create_book(path, names):
    wb = Workbook()
    ws = wb.active
    ws.title = "Сотрудники"
    ws.append(["ФИО", "Email", "Примечание"])
    for name in names:
        ws.append([name, "", ""])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.border = BORDER
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 30
    ws.freeze_panes = "A2"
    wb.save(path)


def load_book(path):
    book = {}
    for row in read_table(path):
        row = {simple(k): (v or "").strip() for k, v in row.items() if k}
        fio = row.get("фио") or row.get("сотрудник", "")
        email = row.get("email") or row.get("e-mail") or row.get("почта", "")
        if fio:
            book[simple_fio(fio)] = email
    return book


def add_emails(people, problems, path=None):
    """Проставляет email сотрудникам из справочника.

    Если справочника ещё нет — создаём его со всеми ФИО из ведомости.
    Если в ведомости появился новый человек — дописываем его в конец справочника.
    """
    if path is None:
        existing = [p for p in DEFAULT_FILES if p.exists()]
        path = existing[0] if existing else DEFAULT_FILES[0]
    path = Path(path)

    if not path.exists():
        create_book(path, [p.fio for p in people])
        problems.append(f"создан справочник {path.name} — впишите в него email сотрудников")
        for p in people:
            p.skip = "нет email"
        return

    book = load_book(path)
    new_people = []
    for p in people:
        if simple_fio(p.fio) not in book:
            p.skip = f"нет в {path.name}"
            new_people.append(p.fio)
            continue
        p.email = book[simple_fio(p.fio)]
        if not p.email:
            p.skip = "нет email"
        elif not looks_like_email(p.email):
            p.skip = f"неверный email «{p.email}»"

    if not new_people:
        return
    if path.suffix.lower() != ".xlsx":
        problems.append(f"нет в {path.name}: {', '.join(new_people)}")
        return
    try:
        wb = load_workbook(path)
        for fio in new_people:
            wb.worksheets[0].append([fio, "", ""])
        wb.save(path)
        problems.append(f"новые сотрудники дописаны в {path.name}, впишите их email: {', '.join(new_people)}")
    except PermissionError:
        problems.append(f"в {path.name} нет: {', '.join(new_people)} (файл открыт в Excel, дописать не получилось)")
